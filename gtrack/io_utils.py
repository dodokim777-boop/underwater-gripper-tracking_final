# -*- coding: utf-8 -*-
"""파일 입출력 공용 함수 (NPZ 값 읽기, 엑셀 서식)"""

import numpy as np


def _npz_value(data, keys, default=None):
    for key in keys:
        if key in data.files:
            value = data[key]
            if np.asarray(value).shape == ():
                return np.asarray(value).item()
            return value
    return default


def format_workbook_sheets(workbook, wide_text_columns=('Description', 'Value', 'Criterion', 'Note')):
    """모든 시트: 첫 행 고정, 머리글 굵게, 열 너비 자동 조정."""
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    for ws in workbook.worksheets:
        ws.freeze_panes = 'A2'
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col_idx, col_cells in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), start=1):
            header = str(col_cells[0].value) if col_cells[0].value is not None else ''
            longest = max([len(header)] + [len(str(c.value)) for c in col_cells[1:] if c.value is not None])
            letter = get_column_letter(col_idx)
            if header in wide_text_columns:
                ws.column_dimensions[letter].width = 90
                for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                    for c in row:
                        c.alignment = Alignment(wrap_text=True, vertical='top')
            else:
                ws.column_dimensions[letter].width = min(max(longest, 6) + 2, 40)
